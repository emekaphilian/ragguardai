from __future__ import annotations

from io import BytesIO
from pathlib import Path


TEXT_EXTENSIONS = {
    ".c",
    ".cfg",
    ".conf",
    ".css",
    ".csv",
    ".html",
    ".htm",
    ".ini",
    ".java",
    ".js",
    ".json",
    ".log",
    ".md",
    ".markdown",
    ".py",
    ".rst",
    ".sql",
    ".svg",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}


class UnsupportedDocumentType(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix in TEXT_EXTENSIONS:
        return _decode_text(content)
    if suffix == ".pdf":
        return _extract_pdf(content)
    if suffix == ".docx":
        return _extract_docx(content)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(content)
    if suffix == ".pptx":
        return _extract_pptx(content)

    raise UnsupportedDocumentType(
        f"Unsupported document type: {suffix or 'file without extension'}"
    )


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        rows.append(f"[{sheet.title}]")
        rows.extend(", ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
    return "\n".join(rows)


def _extract_pptx(content: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if text:
            slides.append(f"[Slide {index}]\n" + "\n".join(text))
    return "\n".join(slides)
